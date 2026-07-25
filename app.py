import json
import logging
import os
import re
import uuid
from datetime import datetime, timedelta
from functools import wraps
from html import escape
from io import BytesIO
from pathlib import Path

from flask import Flask, flash, jsonify, redirect, render_template, request, send_file, send_from_directory, session, url_for
from werkzeug.exceptions import HTTPException
from werkzeug.utils import secure_filename
from werkzeug.security import check_password_hash

from utils.database import (
    delete_interview,
    delete_aptitude_test,
    delete_user,
    complete_onboarding,
    create_notification,
    create_user,
    get_admin_analytics,
    get_dashboard_stats,
    get_default_user_id,
    get_aptitude_test,
    get_interview_analytics,
    get_interview_history,
    get_interview_report,
    get_latest_analysis,
    get_latest_resume,
    get_recent_aptitude_tests,
    get_profile_summary,
    get_profile,
    get_notifications,
    get_settings,
    get_user_by_email,
    get_user_by_id,
    record_user_login,
    global_search,
    list_resumes,
    list_job_applications,
    save_job_application,
    update_job_application,
    delete_job_application,
    get_recent_analyses,
    get_recent_generated_resumes,
    get_recent_interviews,
    init_db,
    save_analysis,
    save_aptitude_test,
    save_generated_resume,
    save_interview,
    save_interview_session,
    update_profile,
    update_resume_state,
    update_settings,
    update_interview_session,
)
from utils.resume_pdf import create_resume_pdf

try:
    from dotenv import load_dotenv
except ModuleNotFoundError:
    load_dotenv = None


BASE_DIR = Path(__file__).resolve().parent
UPLOAD_FOLDER = BASE_DIR / "uploads"
ALLOWED_EXTENSIONS = {"pdf"}
ALLOWED_PROFILE_IMAGE_EXTENSIONS = {"jpg", "jpeg", "png", "gif", "webp"}
MAX_UPLOAD_BYTES = 8 * 1024 * 1024
AI_ANALYSIS_UNAVAILABLE = "AI analysis is temporarily unavailable because the API usage limit has been reached."
GENERIC_ANALYSIS_ERROR = "Resume analysis is temporarily unavailable. An offline analysis has been generated instead."
GENERIC_AI_ASSIST_ERROR = "AI writing assistance is temporarily unavailable. Please try again later."
LOCAL_INTERVIEW_START_MESSAGE = "Interview questions loaded from the local ResumeIQ+ question bank."
LOCAL_APTITUDE_START_MESSAGE = "Aptitude questions loaded from the local ResumeIQ+ question bank."
QUOTA_INTERVIEW_MESSAGE = (
    "Your interview has been completed successfully.\n\n"
    "AI evaluation is temporarily unavailable because today's API limit has been reached.\n\n"
    "Your interview has been saved.\n\n"
    "Click 'Generate AI Feedback' later when the API becomes available."
)

if load_dotenv:
    load_dotenv(BASE_DIR / ".env", override=False)

app = Flask(__name__)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", os.urandom(24))
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.getenv("SESSION_COOKIE_SECURE", "0") == "1"
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(minutes=int(os.getenv("SESSION_TIMEOUT_MINUTES", "120")))
logging.basicConfig(level=logging.INFO)

UPLOAD_FOLDER.mkdir(exist_ok=True)
init_db()


RATE_LIMITS = {}
EMAIL_PATTERN = re.compile(r"^[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}$", re.IGNORECASE)


def current_user():
    user_id = session.get("user_id")
    if user_id:
        return get_user_by_id(user_id)
    return None


def current_user_id():
    user = current_user()
    return user["id"] if user else get_default_user_id()


def wants_json_response():
    if request.is_json:
        return True
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return True
    best = request.accept_mimetypes.best_match(["application/json", "text/html"])
    return best == "application/json" and request.accept_mimetypes[best] > request.accept_mimetypes["text/html"]


def auth_failure(message, status=400, actions=None):
    payload = {"success": False, "message": message, "error": message}
    if actions:
        payload["actions"] = actions
    return jsonify(payload), status


def auth_success(message, redirect_endpoint="dashboard_page", status=200, **values):
    redirect_url = url_for(redirect_endpoint, **values)
    return jsonify({"success": True, "message": message, "redirect": redirect_url, "csrf_token": session.get("csrf_token")}), status


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not current_user():
            if request.path.startswith(("/dashboard-data", "/api/", "/analyze", "/ai-assist", "/interview", "/aptitude", "/resume", "/export")):
                return jsonify({"success": False, "error": "Please log in to continue."}), 401
            return redirect(url_for("login_page"))
        return view(*args, **kwargs)

    return wrapped


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        user = current_user()
        if not user or not user.get("is_admin"):
            return jsonify({"success": False, "error": "Admin access required."}), 403
        return view(*args, **kwargs)

    return wrapped


def csrf_token():
    token = session.get("csrf_token")
    if not token:
        token = uuid.uuid4().hex
        session["csrf_token"] = token
    return token


@app.context_processor
def inject_auth_context():
    return {"auth_user": current_user(), "csrf_token": csrf_token()}


@app.before_request
def protect_state_changes():
    if request.method in {"POST", "PUT", "PATCH", "DELETE"} and not request.path.startswith("/auth/"):
        sent = request.headers.get("X-CSRF-Token") or request.form.get("csrf_token")
        if sent != session.get("csrf_token"):
            return jsonify({"success": False, "error": "Security check failed. Refresh and try again."}), 400


def rate_limited(key, limit=12, window_seconds=60):
    now = datetime.utcnow()
    bucket = [stamp for stamp in RATE_LIMITS.get(key, []) if (now - stamp).total_seconds() < window_seconds]
    if len(bucket) >= limit:
        RATE_LIMITS[key] = bucket
        return True
    bucket.append(now)
    RATE_LIMITS[key] = bucket
    return False


def strong_password_errors(password):
    errors = []
    if len(password) < 8:
        errors.append("at least 8 characters")
    if not any(char.isupper() for char in password):
        errors.append("one uppercase letter")
    if not any(char.islower() for char in password):
        errors.append("one lowercase letter")
    if not any(char.isdigit() for char in password):
        errors.append("one number")
    if not any(not char.isalnum() for char in password):
        errors.append("one symbol")
    return errors


def is_valid_email(email):
    return bool(EMAIL_PATTERN.match(email or ""))


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def allowed_profile_image(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_PROFILE_IMAGE_EXTENSIONS


def looks_like_ai_quota_error(error):
    text = str(error).lower()
    return "429" in text or "resource_exhausted" in text or "quota" in text or "usage limit" in text or "temporarily unavailable" in text


def validate_pdf_upload(file_storage, label):
    if not file_storage or file_storage.filename == "":
        raise ValueError(f"Please upload {label}.")

    if not allowed_file(file_storage.filename):
        raise ValueError(f"{label} must be a PDF file.")

    file_storage.stream.seek(0, os.SEEK_END)
    size = file_storage.stream.tell()
    file_storage.stream.seek(0)

    if size <= 0:
        raise ValueError(f"{label} is empty.")

    if size > MAX_UPLOAD_BYTES:
        raise ValueError(f"{label} must be smaller than 8 MB.")

    header = file_storage.stream.read(5)
    file_storage.stream.seek(0)
    if header != b"%PDF-":
        raise ValueError(f"{label} does not look like a valid PDF.")


def save_uploaded_pdf(file_storage):
    original_name = secure_filename(file_storage.filename or "resume.pdf")
    stem = Path(original_name).stem[:80] or "resume"
    filename = f"{stem}-{uuid.uuid4().hex[:10]}.pdf"
    file_path = app.config["UPLOAD_FOLDER"] / filename
    file_storage.save(file_path)
    return filename, file_path


def build_export_payload(job_role, result, comparison=None):
    reports = [report for report in [result, comparison] if report]
    return {"job_role": job_role, "reports": reports}


def as_list(value):
    if isinstance(value, list):
        return [str(item) for item in value if str(item).strip()]

    if isinstance(value, str) and value.strip():
        return [value]

    return []


def get_sub_score(report, key):
    for item in report.get("sub_scores", []):
        if item.get("key") == key:
            return item.get("score")

    return None


def format_score(score):
    return f"{score}%" if score is not None else "N/A"


def build_pdf_styles():
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("ReportTitle", parent=styles["Title"], alignment=TA_CENTER, fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.HexColor("#111827"), spaceAfter=6))
    styles.add(ParagraphStyle("ReportSubtitle", parent=styles["BodyText"], alignment=TA_CENTER, fontSize=10, leading=14, textColor=colors.HexColor("#4B5563"), spaceAfter=12))
    styles.add(ParagraphStyle("MetaText", parent=styles["BodyText"], alignment=TA_CENTER, fontSize=9, leading=12, textColor=colors.HexColor("#374151"), spaceAfter=4))
    styles.add(ParagraphStyle("SectionHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=colors.HexColor("#0F766E"), spaceBefore=12, spaceAfter=8))
    styles.add(ParagraphStyle("ResumeHeading", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=colors.HexColor("#111827"), spaceBefore=14, spaceAfter=8))
    styles.add(ParagraphStyle("BodyClean", parent=styles["BodyText"], alignment=TA_LEFT, fontSize=9.5, leading=14, textColor=colors.HexColor("#1F2937"), spaceAfter=5))
    styles.add(ParagraphStyle("Muted", parent=styles["BodyText"], fontSize=9, leading=13, textColor=colors.HexColor("#6B7280"), spaceAfter=5))
    styles.add(ParagraphStyle("ScoreLarge", parent=styles["BodyText"], alignment=TA_CENTER, fontName="Helvetica-Bold", fontSize=26, leading=30, textColor=colors.HexColor("#0F766E")))

    return styles


class NumberedCanvas:
    def __init__(self, canvas, footer_text):
        self.canvas = canvas
        self.footer_text = footer_text
        self.saved_pages = []

    def __getattr__(self, name):
        return getattr(self.canvas, name)

    def showPage(self):
        self.saved_pages.append(dict(self.canvas.__dict__))
        self.canvas._startPage()

    def save(self):
        page_count = len(self.saved_pages)

        for page_number, page_state in enumerate(self.saved_pages, start=1):
            self.canvas.__dict__.update(page_state)
            self.draw_footer(page_number, page_count)
            self.canvas.showPage()

        self.canvas.save()

    def draw_footer(self, page_number, page_count):
        from reportlab.lib import colors

        self.canvas.saveState()
        self.canvas.setStrokeColor(colors.HexColor("#D1D5DB"))
        self.canvas.setLineWidth(0.5)
        self.canvas.line(54, 34, 558, 34)
        self.canvas.setFont("Helvetica", 8)
        self.canvas.setFillColor(colors.HexColor("#6B7280"))
        self.canvas.drawString(54, 22, self.footer_text)
        self.canvas.drawRightString(558, 22, f"Page {page_number} / {page_count}")
        self.canvas.restoreState()


def divider(width=510):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    line = Table([[""]], colWidths=[width], rowHeights=[1])
    line.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#D1D5DB"))]))
    return line


def draw_header(story, styles, job_role):
    from reportlab.platypus import Paragraph, Spacer

    generated_at = datetime.now().strftime("%B %d, %Y at %I:%M %p")

    story.append(Paragraph("ResumeIQ+", styles["ReportTitle"]))
    story.append(Paragraph("AI-Powered Resume Analysis Report", styles["ReportSubtitle"]))
    story.append(Paragraph(f"<b>Target Job Role:</b> {escape(str(job_role or 'Not provided'))}", styles["MetaText"]))
    story.append(Paragraph(f"<b>Generated:</b> {generated_at}", styles["MetaText"]))
    story.append(Spacer(1, 8))
    story.append(divider())
    story.append(Spacer(1, 14))


def draw_section_heading(story, styles, title):
    from reportlab.platypus import Paragraph, Spacer

    story.append(Paragraph(escape(title), styles["SectionHeading"]))
    story.append(divider())
    story.append(Spacer(1, 9))


def draw_text_section(story, styles, title, text, fallback):
    from reportlab.platypus import Paragraph, Spacer

    draw_section_heading(story, styles, title)
    story.append(Paragraph(escape(str(text or fallback)), styles["BodyClean"]))
    story.append(Spacer(1, 8))


def draw_summary_box(story, styles, text):
    from reportlab.lib import colors
    from reportlab.platypus import KeepTogether, Paragraph, Spacer, Table, TableStyle

    summary = escape(str(text or "No resume summary available."))
    box = Table([[Paragraph(summary, styles["BodyClean"])]], colWidths=[480])
    box.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F9FAFB")),
                ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#D1D5DB")),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ]
        )
    )
    story.append(KeepTogether([Paragraph("Executive Summary", styles["SectionHeading"]), divider(), Spacer(1, 9), box]))
    story.append(Spacer(1, 10))


def draw_bullet_section(story, styles, title, items, fallback):
    from reportlab.platypus import KeepTogether, Paragraph, Spacer

    content = [Paragraph(escape(title), styles["SectionHeading"]), divider(), Spacer(1, 9)]
    items = as_list(items)

    if not items:
        content.append(Paragraph(escape(fallback), styles["Muted"]))
    else:
        for item in items:
            content.append(Paragraph(f"&bull; {escape(str(item))}", styles["BodyClean"]))

    story.append(KeepTogether(content))
    story.append(Spacer(1, 8))


def draw_ats_card(story, styles, report):
    from reportlab.lib import colors
    from reportlab.platypus import KeepTogether, Paragraph, Spacer, Table, TableStyle

    score = format_score(report.get("ats_score"))
    data = [
        [Paragraph("Overall ATS Score", styles["BodyClean"]), Paragraph(score, styles["ScoreLarge"])],
        ["Skills Match", format_score(get_sub_score(report, "skills_match"))],
        ["Experience Match", format_score(get_sub_score(report, "experience_match"))],
        ["Resume Formatting", format_score(get_sub_score(report, "resume_formatting"))],
        ["Keyword Optimization", format_score(get_sub_score(report, "keyword_optimization"))],
    ]

    content = [Paragraph("ATS Metrics", styles["SectionHeading"]), divider(), Spacer(1, 9)]
    table = Table(data, colWidths=[310, 170])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#CCFBF1")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0F766E")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 15),
                ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                ("ALIGN", (1, 0), (1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
                ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#99F6E4")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    content.append(table)
    story.append(KeepTogether(content))
    story.append(Spacer(1, 12))


def draw_keywords(story, styles, report):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    keywords = as_list(report.get("important_keywords"))
    if not keywords:
        return

    draw_section_heading(story, styles, "Skills Match")
    rows = []
    row = []

    for keyword in keywords:
        row.append(Paragraph(escape(keyword), styles["BodyClean"]))
        if len(row) == 3:
            rows.append(row)
            row = []

    if row:
        while len(row) < 3:
            row.append("")
        rows.append(row)

    table = Table(rows, colWidths=[160, 160, 160])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#ECFDF5")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#A7F3D0")),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1FAE5")),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 8))


def categorize_keywords(keywords):
    categories = {
        "Frontend": [],
        "Backend": [],
        "Tools": [],
        "Soft Skills": [],
    }
    frontend_terms = {"html", "css", "javascript", "react", "vue", "angular", "ui", "ux", "frontend", "tailwind"}
    backend_terms = {"python", "java", "node", "flask", "django", "api", "sql", "database", "backend", "server"}
    tool_terms = {"git", "docker", "aws", "azure", "gcp", "tableau", "power bi", "excel", "kubernetes", "ci/cd"}

    for keyword in keywords:
        lowered = keyword.lower()
        if any(term in lowered for term in frontend_terms):
            categories["Frontend"].append(keyword)
        elif any(term in lowered for term in backend_terms):
            categories["Backend"].append(keyword)
        elif any(term in lowered for term in tool_terms):
            categories["Tools"].append(keyword)
        else:
            categories["Soft Skills"].append(keyword)

    return categories


def draw_compact_skills(story, styles, report):
    from reportlab.lib import colors
    from reportlab.platypus import KeepTogether, Paragraph, Spacer, Table, TableStyle

    keywords = as_list(report.get("important_keywords"))
    if not keywords:
        keywords = as_list(report.get("strengths"))

    if not keywords:
        return

    categories = categorize_keywords(keywords)
    data = [["Category", "Skills / Keywords"]]

    for category, values in categories.items():
        if values:
            data.append([category, ", ".join(values[:8])])

    if len(data) == 1:
        return

    table = Table(data, colWidths=[130, 350])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111827")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.append(KeepTogether([Paragraph("Compact Skills View", styles["SectionHeading"]), divider(), Spacer(1, 9), table]))
    story.append(Spacer(1, 10))


def draw_priority_recommendations(story, styles, report):
    from reportlab.lib import colors
    from reportlab.platypus import KeepTogether, Paragraph, Spacer, Table, TableStyle

    suggestions = as_list(report.get("improvement_suggestions"))
    missing = as_list(report.get("missing_skills"))
    priorities = suggestions[:2] or missing[:2]

    if not priorities:
        priorities = ["No urgent priority improvements were identified."]

    body = [Paragraph(f"&bull; {escape(item)}", styles["BodyClean"]) for item in priorities]
    box = Table([[body]], colWidths=[480])
    box.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFFBEB")),
                ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#FCD34D")),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ]
        )
    )
    story.append(KeepTogether([Paragraph("Top Priority Improvements", styles["SectionHeading"]), divider(), Spacer(1, 9), box]))
    story.append(Spacer(1, 10))


def draw_comparison_table(story, styles, reports):
    from reportlab.lib import colors
    from reportlab.platypus import KeepTogether, Paragraph, Spacer, Table, TableStyle

    if len(reports) < 2:
        return

    resume_a, resume_b = reports[0], reports[1]
    data = [
        ["Metric", resume_a.get("label", "Resume A"), resume_b.get("label", "Resume B")],
        ["ATS Score", format_score(resume_a.get("ats_score")), format_score(resume_b.get("ats_score"))],
        ["Skills Match", format_score(get_sub_score(resume_a, "skills_match")), format_score(get_sub_score(resume_b, "skills_match"))],
        ["Experience Match", format_score(get_sub_score(resume_a, "experience_match")), format_score(get_sub_score(resume_b, "experience_match"))],
        ["Formatting", format_score(get_sub_score(resume_a, "resume_formatting")), format_score(get_sub_score(resume_b, "resume_formatting"))],
        ["Keyword Optimization", format_score(get_sub_score(resume_a, "keyword_optimization")), format_score(get_sub_score(resume_b, "keyword_optimization"))],
    ]

    table = Table(data, colWidths=[160, 160, 160])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111827")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9CA3AF")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 9),
                ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(KeepTogether([Paragraph("Resume Comparison", styles["SectionHeading"]), divider(), Spacer(1, 9), table]))
    story.append(Spacer(1, 14))


def draw_single_report(story, styles, report, add_page_break=False):
    from reportlab.platypus import PageBreak, Paragraph, Spacer

    if add_page_break:
        story.append(PageBreak())

    story.append(Paragraph(escape(str(report.get("label", "Resume"))), styles["ResumeHeading"]))
    story.append(Spacer(1, 4))
    draw_summary_box(story, styles, report.get("resume_summary"))
    draw_ats_card(story, styles, report)
    draw_priority_recommendations(story, styles, report)
    draw_keywords(story, styles, report)
    draw_compact_skills(story, styles, report)
    draw_bullet_section(story, styles, "Strengths", report.get("strengths"), "No strengths available.")
    draw_bullet_section(story, styles, "Missing Skills", report.get("missing_skills"), "No missing skills available.")
    draw_bullet_section(story, styles, "Improvement Suggestions", report.get("improvement_suggestions"), "No improvement suggestions available.")
    draw_bullet_section(story, styles, "Recommended Projects", report.get("recommended_projects"), "No recommended projects available.")


def create_report_pdf(payload):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    from reportlab.platypus import SimpleDocTemplate

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        title="ResumeIQ+ Report",
        leftMargin=54,
        rightMargin=54,
        topMargin=48,
        bottomMargin=48,
    )
    styles = build_pdf_styles()
    story = []
    job_role = payload.get("job_role", "Target Role")
    reports = [report for report in payload.get("reports", []) if isinstance(report, dict)]

    draw_header(story, styles, job_role)
    draw_comparison_table(story, styles, reports)

    if not reports:
        draw_text_section(story, styles, "Executive Summary", None, "No analysis data available.")
    else:
        for index, report in enumerate(reports):
            draw_single_report(story, styles, report, add_page_break=index > 0)

    def make_canvas(*args, **kwargs):
        return NumberedCanvas(canvas.Canvas(*args, **kwargs), "ResumeIQ+ | AI Resume Analysis | Generated Automatically")

    doc.build(story, canvasmaker=make_canvas)
    buffer.seek(0)
    return buffer


def create_interview_pdf(report):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    setup = report.get("setup") or {}
    scores = report.get("scores") or {}
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, title="ResumeIQ+ Interview Report", leftMargin=48, rightMargin=48, topMargin=42, bottomMargin=42)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("InterviewTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=21, leading=25, textColor=colors.HexColor("#111827"), spaceAfter=6))
    styles.add(ParagraphStyle("InterviewMeta", parent=styles["BodyText"], fontSize=9.5, leading=13, textColor=colors.HexColor("#4B5563"), spaceAfter=4))
    styles.add(ParagraphStyle("InterviewHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#0F766E"), spaceBefore=12, spaceAfter=6))
    styles.add(ParagraphStyle("InterviewBody", parent=styles["BodyText"], fontSize=9, leading=13, textColor=colors.HexColor("#1F2937"), spaceAfter=5))

    story = [
        Paragraph("ResumeIQ+ Interview Coach Report", styles["InterviewTitle"]),
        Paragraph(f"<b>Candidate:</b> {escape(str(setup.get('candidate_name') or 'Candidate'))}", styles["InterviewMeta"]),
        Paragraph(f"<b>Role:</b> {escape(str(setup.get('job_role') or 'General Role'))}", styles["InterviewMeta"]),
        Paragraph(f"<b>Interview:</b> {escape(str(setup.get('interview_type') or 'Mixed'))} | {escape(str(setup.get('difficulty') or 'Medium'))} | {escape(str(setup.get('experience_level') or 'Entry'))}", styles["InterviewMeta"]),
        Paragraph(f"<b>Questions:</b> {escape(str(setup.get('question_count') or len(report.get('answers') or [])))}", styles["InterviewMeta"]),
        Spacer(1, 10),
    ]
    score_rows = [["Metric", "Score"]]
    for key, label in [
        ("technical_score", "Technical"),
        ("communication_score", "Communication"),
        ("confidence_score", "Confidence"),
        ("grammar_score", "Grammar"),
        ("vocabulary_score", "Vocabulary"),
        ("problem_solving_score", "Problem Solving"),
        ("overall_score", "Overall"),
    ]:
        score_rows.append([label, f"{scores.get(key, 0) or 0}%"])
    score_table = Table(score_rows, colWidths=[300, 120])
    score_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111827")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        ("ALIGN", (1, 1), (-1, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([Paragraph("Scores", styles["InterviewHeading"]), score_table, Spacer(1, 8)])
    story.append(Paragraph("Performance Summary", styles["InterviewHeading"]))
    story.append(Paragraph(escape(str(report.get("performance_summary") or "No summary available.")), styles["InterviewBody"]))
    story.append(Paragraph(f"<b>Hiring Recommendation:</b> {escape(str(report.get('hiring_recommendation') or 'Needs more preparation'))}", styles["InterviewBody"]))
    for title, key in [("Strengths", "strengths"), ("Weaknesses", "weaknesses"), ("Improvement Plan", "improvement_plan"), ("Learning Resources", "learning_resources"), ("Suggested Courses", "suggested_courses")]:
        story.append(Paragraph(title, styles["InterviewHeading"]))
        items = as_list(report.get(key))
        if not items:
            story.append(Paragraph("No items available.", styles["InterviewBody"]))
        for item in items:
            story.append(Paragraph(f"&bull; {escape(str(item))}", styles["InterviewBody"]))
    story.append(Paragraph("Questions and Answers", styles["InterviewHeading"]))
    for index, item in enumerate(report.get("answers") or [], start=1):
        story.append(Paragraph(f"<b>Q{index}:</b> {escape(str(item.get('question') or ''))}", styles["InterviewBody"]))
        story.append(Paragraph(f"<b>Answer:</b> {escape(str(item.get('answer') or 'Skipped'))}", styles["InterviewBody"]))
        story.append(Paragraph(f"<b>Score:</b> {item.get('score', 0)}% | <b>Better answer:</b> {escape(str(item.get('better_answer') or ''))}", styles["InterviewBody"]))
        for tip in as_list(item.get("improvement_tips")):
            story.append(Paragraph(f"&bull; {escape(str(tip))}", styles["InterviewBody"]))
        story.append(Spacer(1, 4))
    doc.build(story)
    buffer.seek(0)
    return buffer


def create_aptitude_pdf(result):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    setup = result.get("setup") or {}
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, title="ResumeIQ+ Aptitude Report", leftMargin=48, rightMargin=48, topMargin=42, bottomMargin=42)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("AptTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=21, leading=25, textColor=colors.HexColor("#111827"), spaceAfter=8))
    styles.add(ParagraphStyle("AptBody", parent=styles["BodyText"], fontSize=9, leading=13, textColor=colors.HexColor("#1F2937"), spaceAfter=5))
    styles.add(ParagraphStyle("AptHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#0F766E"), spaceBefore=12, spaceAfter=6))
    story = [
        Paragraph("ResumeIQ+ Aptitude Test Report", styles["AptTitle"]),
        Paragraph(f"<b>Candidate:</b> {escape(str(setup.get('candidate_name') or 'Candidate'))}", styles["AptBody"]),
        Paragraph(f"<b>Category:</b> {escape(str(setup.get('category') or 'Mixed'))} | <b>Difficulty:</b> {escape(str(setup.get('difficulty') or 'Medium'))}", styles["AptBody"]),
        Spacer(1, 8),
    ]
    rows = [
        ["Correct", str(result.get("correct_answers", 0))],
        ["Wrong", str(result.get("wrong_answers", 0))],
        ["Accuracy", f"{result.get('accuracy', 0)}%"],
        ["Final Score", f"{result.get('final_score', 0)}%"],
    ]
    table = Table(rows, colWidths=[220, 120])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F9FAFB")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([Paragraph("Summary", styles["AptHeading"]), table])
    for title, key in [("Weak Areas", "weak_areas"), ("Strong Areas", "strong_areas")]:
        story.append(Paragraph(title, styles["AptHeading"]))
        for item in as_list(result.get(key)):
            story.append(Paragraph(f"&bull; {escape(str(item))}", styles["AptBody"]))
    story.append(Paragraph("Question Review", styles["AptHeading"]))
    for index, item in enumerate(result.get("questions") or [], start=1):
        selected = item.get("selected_answer") or "Not answered"
        correct = item.get("correct_answer") or ""
        status = "Correct" if item.get("is_correct") else "Wrong"
        story.append(Paragraph(f"<b>Q{index}:</b> {escape(str(item.get('question') or ''))}", styles["AptBody"]))
        story.append(Paragraph(f"Selected: {escape(str(selected))} | Correct: {escape(str(correct))} | {status}", styles["AptBody"]))
        story.append(Paragraph(escape(str(item.get("explanation") or "")), styles["AptBody"]))
    doc.build(story)
    buffer.seek(0)
    return buffer


def create_resume_docx(resume):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ModuleNotFoundError as exc:
        raise ValueError("DOCX export requires python-docx. Run pip install -r requirements.txt.") from exc

    buffer = BytesIO()
    document = Document()
    styles = document.styles
    font_name = str(resume.get("font") or "Arial")
    styles["Normal"].font.name = font_name
    styles["Normal"].font.size = Pt(10)
    accent = str(resume.get("accent_color") or resume.get("color_theme") or "#0F766E").strip().lstrip("#")
    if len(accent) != 6:
        accent = "0F766E"
    accent_rgb = RGBColor(int(accent[0:2], 16), int(accent[2:4], 16), int(accent[4:6], 16))
    personal = resume.get("personal_details") or {}
    name = personal.get("full_name") or "Your Name"
    if resume.get("show_photo"):
        photo = document.add_paragraph()
        run = photo.add_run((name[:1] or "R").upper())
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = accent_rgb
    heading = document.add_heading(str(name), 0)
    heading.runs[0].font.color.rgb = accent_rgb
    contact_bits = [
        personal.get("email"),
        personal.get("phone"),
        personal.get("location"),
        personal.get("portfolio"),
        personal.get("linkedin"),
    ]
    contact = " | ".join(str(bit).strip() for bit in contact_bits if str(bit or "").strip())
    if contact:
        document.add_paragraph(contact)

    def add_section(title, value):
        items = as_list(value)
        if not items:
            return
        section_heading = document.add_heading(title, level=2)
        if section_heading.runs:
            section_heading.runs[0].font.color.rgb = accent_rgb
        for item in items:
            document.add_paragraph(str(item), style="List Bullet")

    section_titles = dict([
        ("education", "Education"),
        ("experience", "Experience"),
        ("projects", "Projects"),
        ("internships", "Internships"),
        ("skills", "Skills"),
        ("technical_skills", "Technical Skills"),
        ("soft_skills", "Soft Skills"),
        ("achievements", "Achievements"),
        ("certifications", "Certifications"),
        ("languages", "Languages"),
        ("interests", "Interests"),
        ("references", "References"),
        ("social_links", "Social Links"),
    ])
    section_titles["career_objective"] = "Career Objective"
    order = resume.get("section_order") if isinstance(resume.get("section_order"), list) else list(section_titles.keys())
    for key in order:
        title = section_titles.get(key)
        if not title:
            continue
        if key == "career_objective":
            objective = str(resume.get("career_objective") or "").strip()
            if objective:
                objective_heading = document.add_heading("Career Objective", level=2)
                if objective_heading.runs:
                    objective_heading.runs[0].font.color.rgb = accent_rgb
                document.add_paragraph(objective)
            continue
        add_section(title, resume.get(key))

    document.save(buffer)
    buffer.seek(0)
    return buffer


@app.route("/")
def index():
    if current_user():
        return redirect(url_for("dashboard_page"))
    return render_template("index.html")


@app.route("/dashboard")
@login_required
def dashboard_page():
    return render_template("index.html", initial_view="home")


@app.route("/login")
def login_page():
    if current_user():
        return redirect(url_for("dashboard_page"))
    return render_template("index.html", auth_mode="login")


@app.route("/signup")
def signup_page():
    if current_user():
        return redirect(url_for("dashboard_page"))
    return render_template("index.html", auth_mode="signup")


@app.route("/forgot-password")
def forgot_password_page():
    return render_template("index.html", auth_mode="forgot")


@app.route("/reset-password")
def reset_password_page():
    return render_template("index.html", auth_mode="reset")


@app.route("/verify-email")
def verify_email_page():
    return render_template("index.html", auth_mode="verify")


@app.route("/logout")
def logout():
    session.clear()
    flash("You have been logged out.", "success")
    return redirect(url_for("login_page"))


@app.route("/uploads/<path:filename>")
@login_required
def uploaded_file(filename):
    # Profile images are user-namespaced; do not expose another account's upload by URL.
    if not filename.startswith(f"profile-{current_user_id()}-"):
        return jsonify({"success": False, "error": "File not found."}), 404
    return send_from_directory(app.config["UPLOAD_FOLDER"], filename)


@app.route("/auth/signup", methods=["POST"])
def auth_signup():
    payload = request.get_json(silent=True) or request.form
    email = str(payload.get("email") or "").strip().lower()
    full_name = str(payload.get("full_name") or "").strip()
    password = str(payload.get("password") or "")
    confirm_password = str(payload.get("confirm_password") or payload.get("confirmPassword") or "")
    if rate_limited(f"signup:{request.remote_addr}", 8, 300):
        return auth_failure("Too many sign-up attempts. Try again shortly.", 429)
    if not full_name or not is_valid_email(email):
        return auth_failure("Invalid email or missing name.", 400)
    errors = strong_password_errors(password)
    if errors:
        return auth_failure("Password needs " + ", ".join(errors) + ".", 400)
    if confirm_password != password:
        return auth_failure("Passwords do not match.", 400)
    if get_user_by_email(email):
        return auth_failure(
            "Email already registered.",
            409,
            actions=[
                {"label": "Go to Login", "href": url_for("login_page")},
                {"label": "Forgot Password", "href": url_for("forgot_password_page")},
            ],
        )
    try:
        create_user(email, password, full_name)
    except Exception as error:
        # The database UNIQUE constraint is the final safeguard against concurrent sign-ups.
        if "unique" in str(error).lower():
            return auth_failure("Email already registered.", 409)
        app.logger.exception("User signup failed")
        return auth_failure("Server unavailable. Please try again shortly.", 500)
    session.clear()
    session["csrf_token"] = uuid.uuid4().hex
    return auth_success("Account created successfully.", "login_page", status=201)


@app.route("/auth/login", methods=["POST"])
def auth_login():
    payload = request.get_json(silent=True) or request.form
    email = str(payload.get("email") or "").strip().lower()
    password = str(payload.get("password") or "")
    if not is_valid_email(email):
        return auth_failure("Invalid email.", 400)
    if rate_limited(f"login:{request.remote_addr}:{email}", 10, 300):
        return auth_failure("Too many login attempts. Try again shortly.", 429)
    user = get_user_by_email(email)
    if not user or not check_password_hash(user.get("password_hash", ""), password):
        return auth_failure("Invalid email or password.", 401)
    session.clear()
    session["user_id"] = user["id"]
    session["user_name"] = user["full_name"]
    session["user_email"] = user["email"]
    session["csrf_token"] = uuid.uuid4().hex
    session.permanent = bool(payload.get("remember"))
    record_user_login(user["id"])
    return auth_success("Welcome back.")


@app.route("/auth/forgot-password", methods=["POST"])
def auth_forgot_password():
    payload = request.get_json(silent=True) or request.form
    email = str(payload.get("email") or "").strip().lower()
    if not is_valid_email(email):
        return jsonify({"success": False, "error": "Enter a valid email address."}), 400
    if get_user_by_email(email):
        app.logger.info("Password reset requested for %s. Demo reset page: /reset-password", email)
    return jsonify({"success": True, "message": "If that email exists, a reset link has been generated for this demo environment."})


@app.route("/auth/reset-password", methods=["POST"])
def auth_reset_password():
    payload = request.get_json(silent=True) or request.form
    password = str(payload.get("password") or "")
    errors = strong_password_errors(password)
    if errors:
        return jsonify({"success": False, "error": "Password needs " + ", ".join(errors) + "."}), 400
    return jsonify({"success": True, "message": "Password reset flow is ready for an email-token provider in production."})


@app.route("/auth/verify-email", methods=["POST"])
def auth_verify_email():
    if current_user():
        create_notification(current_user_id(), "Email verification requested", "Verification is marked ready for production email delivery.", "security")
    return jsonify({"success": True, "message": "Email verification flow is ready for production email delivery."})


@app.route("/dashboard-data")
@login_required
def dashboard_data():
    user_id = current_user_id()
    return jsonify(
        {
            "success": True,
            "stats": get_dashboard_stats(user_id),
            "recent": get_recent_analyses(user_id=user_id),
            "recent_resumes": get_recent_generated_resumes(user_id=user_id),
            "recent_interviews": get_recent_interviews(user_id=user_id),
            "recent_aptitude_tests": get_recent_aptitude_tests(user_id=user_id),
            "interview_analytics": get_interview_analytics(user_id),
            "latest_resume": get_latest_resume(user_id),
            "latest_analysis": get_latest_analysis(user_id),
            "profile": get_profile_summary(user_id),
            "notifications": get_notifications(user_id),
            "settings": get_settings(user_id),
        }
    )


@app.route("/analyze", methods=["POST"])
@login_required
def analyze():
    result = None
    comparison = None
    offline_used = False

    resume = request.files.get("resume")
    compare_resume = request.files.get("compare_resume")
    job_role = request.form.get("job_role", "").strip()

    try:
        validate_pdf_upload(resume, "a PDF resume")
        if compare_resume and compare_resume.filename:
            validate_pdf_upload(compare_resume, "the comparison resume")

        if not job_role:
            return jsonify({"success": False, "error": "Please enter a target job role."}), 400

        filename, file_path = save_uploaded_pdf(resume)

        from utils.analyzer import analyze_resume, locally_analyze_resume

        def run_resume_analysis(path, label):
            nonlocal offline_used
            try:
                report = analyze_resume(path, job_role)
            except Exception as exc:
                app.logger.warning("AI resume analysis unavailable for %s: %s", label, exc, exc_info=True)
                offline_used = True
                report = locally_analyze_resume(path, job_role)
                report["offline_message"] = AI_ANALYSIS_UNAVAILABLE if looks_like_ai_quota_error(exc) else GENERIC_ANALYSIS_ERROR
            report["label"] = label
            return report

        app.logger.info("Starting Gemini analysis for primary resume: %s", filename)
        result = run_resume_analysis(file_path, "Resume A")
        save_analysis(filename, job_role, result, current_user_id())

        if compare_resume and compare_resume.filename:
            compare_filename, compare_path = save_uploaded_pdf(compare_resume)

            app.logger.info("Starting Gemini analysis for comparison resume: %s", compare_filename)
            comparison = run_resume_analysis(compare_path, "Resume B")
            save_analysis(compare_filename, job_role, comparison, current_user_id())
    except ValueError as exc:
        app.logger.warning("Resume analysis validation/configuration error: %s", exc, exc_info=True)
        return jsonify({"success": False, "error": str(exc)}), 400
    except Exception as exc:
        app.logger.exception("Unexpected resume analysis failure")
        return jsonify({"success": False, "error": "Resume analysis could not be completed right now. Please try again."}), 500

    export_payload = build_export_payload(job_role, result, comparison)

    return jsonify(
        {
            "success": True,
            "job_role": job_role,
            "result": result,
            "comparison": comparison,
            "export_payload": export_payload,
            "offline": offline_used,
            "message": AI_ANALYSIS_UNAVAILABLE if offline_used else "",
        }
    )


@app.route("/ai-assist", methods=["POST"])
@login_required
def ai_assist():
    payload = request.get_json(silent=True) or {}
    action = str(payload.get("action") or "Improve resume content").strip()
    section = str(payload.get("section") or "Resume").strip()
    content = str(payload.get("content") or "").strip()
    target_role = str(payload.get("target_role") or "").strip()

    try:
        from utils.analyzer import generate_resume_assistance

        result = generate_resume_assistance(action, section, content, target_role)
        return jsonify({"success": True, "result": result})
    except ValueError as exc:
        app.logger.warning("AI builder assistance unavailable: %s", exc, exc_info=True)
        return jsonify({"success": False, "error": GENERIC_AI_ASSIST_ERROR}), 400
    except Exception:
        app.logger.exception("AI builder assistance failed")
        return jsonify({"success": False, "error": GENERIC_AI_ASSIST_ERROR}), 500


@app.route("/interview/start", methods=["POST"])
@login_required
def interview_start():
    payload = request.get_json(silent=True) or {}
    try:
        question_count = int(payload.get("question_count") or 5)
    except (TypeError, ValueError):
        question_count = 5

    setup = {
        "candidate_name": str(payload.get("candidate_name") or "Candidate").strip()[:120],
        "job_role": str(payload.get("job_role") or "").strip()[:160],
        "experience_level": str(payload.get("experience_level") or "Entry").strip()[:80],
        "difficulty": str(payload.get("difficulty") or "Medium").strip()[:80],
        "interview_type": str(payload.get("interview_type") or "Mixed").strip()[:80],
        "question_count": max(1, min(question_count, 20)),
        "custom_focus": str(payload.get("custom_focus") or "").strip()[:500],
    }
    if not setup["job_role"]:
        return jsonify({"success": False, "error": "Please enter a job role for the interview."}), 400

    try:
        from utils.question_bank import get_interview_questions

        questions = get_interview_questions(setup)
        offline = True
        message = LOCAL_INTERVIEW_START_MESSAGE
        session = {
            "setup": setup,
            "questions": questions,
            "answers": ["" for _ in questions],
            "status": "local_active",
        }
        session_id = save_interview_session(session, current_user_id())
        return jsonify({"success": True, "setup": setup, "questions": questions, "session_id": session_id, "offline": offline, "message": message})
    except ValueError as exc:
        app.logger.warning("Interview setup validation failed: %s", exc, exc_info=True)
        return jsonify({"success": False, "error": "Could not start the interview. Please check the setup details and try again."}), 400
    except Exception:
        app.logger.exception("Interview question generation failed")
        return jsonify({"success": False, "error": "Could not start the interview right now. Please try again."}), 500


@app.route("/interview/evaluate-answer", methods=["POST"])
@login_required
def interview_evaluate_answer():
    payload = request.get_json(silent=True) or {}
    answer = str(payload.get("answer") or "").strip()
    words = len(answer.split())
    substance = min(100, max(0, words * 2))
    structure_bonus = 15 if any(token in answer.lower() for token in ["result", "because", "example", "impact", "implemented"]) else 0
    score = min(100, max(25 if answer else 0, substance + structure_bonus))
    evaluation = {
        "technical_score": score,
        "communication_score": min(100, score + 5),
        "confidence_score": min(100, score + 3),
        "grammar_score": min(100, score + 4),
        "problem_solving_score": score,
        "overall_score": score,
        "strengths": ["Answer saved locally."],
        "weaknesses": ["Full AI feedback is generated only after the interview is completed."],
        "better_sample_answer": "",
        "improvement_suggestions": ["Add context, action, result, and measurable impact."],
    }
    return jsonify({"success": True, "evaluation": evaluation, "local": True})


@app.route("/interview/session/<int:session_id>/answers", methods=["POST"])
@login_required
def interview_save_answers(session_id):
    payload = request.get_json(silent=True) or {}
    setup = payload.get("setup") if isinstance(payload.get("setup"), dict) else {}
    questions = payload.get("questions") if isinstance(payload.get("questions"), list) else []
    answers = payload.get("answers") if isinstance(payload.get("answers"), list) else []
    status = str(payload.get("status") or "active").strip()[:40]
    session = {
        "setup": setup,
        "questions": questions,
        "answers": answers,
        "status": status,
    }
    if not update_interview_session(session_id, session, status=status, user_id=current_user_id()):
        return jsonify({"success": False, "error": "Interview session not found."}), 404
    return jsonify({"success": True})


@app.route("/interview/evaluate", methods=["POST"])
@login_required
def interview_evaluate():
    payload = request.get_json(silent=True) or {}
    setup = payload.get("setup") if isinstance(payload.get("setup"), dict) else {}
    answers = payload.get("answers") if isinstance(payload.get("answers"), list) else []
    session_id = payload.get("session_id")
    try:
        from utils.analyzer import GEMINI_QUOTA_MESSAGE, evaluate_interview, locally_evaluate_interview

        report = evaluate_interview(setup, answers)
        report["id"] = save_interview(report, current_user_id())
        if session_id:
            update_interview_session(session_id, {"setup": setup, "answers": answers, "report_id": report["id"], "status": "evaluated"}, status="evaluated", user_id=current_user_id())
        return jsonify({"success": True, "report": report})
    except ValueError as exc:
        app.logger.warning("Interview AI evaluation unavailable: %s", exc, exc_info=True)
        message = QUOTA_INTERVIEW_MESSAGE if looks_like_ai_quota_error(exc) else "AI evaluation is temporarily unavailable. Offline Evaluation is shown immediately."
        report = locally_evaluate_interview(setup, answers, message)
        report["id"] = save_interview(report, current_user_id())
        if session_id:
            update_interview_session(session_id, {"setup": setup, "answers": answers, "report_id": report["id"], "status": "offline_evaluated"}, status="offline_evaluated", user_id=current_user_id())
        return jsonify({"success": True, "report": report, "offline": True, "message": message})
    except Exception as exc:
        app.logger.exception("Interview evaluation failed")
        from utils.analyzer import locally_evaluate_interview

        report = locally_evaluate_interview(setup, answers, "AI evaluation failed. Offline Evaluation is shown immediately.")
        report["id"] = save_interview(report, current_user_id())
        if session_id:
            update_interview_session(session_id, {"setup": setup, "answers": answers, "report_id": report["id"], "status": "offline_evaluated"}, status="offline_evaluated", user_id=current_user_id())
        return jsonify({"success": True, "report": report, "offline": True, "message": "AI evaluation failed. Offline Evaluation is shown immediately."})


@app.route("/interview/history/<int:interview_id>/generate-ai", methods=["POST"])
@login_required
def interview_generate_ai(interview_id):
    report = get_interview_report(interview_id, current_user_id())
    if not report:
        return jsonify({"success": False, "error": "Interview report not found."}), 404
    try:
        from utils.analyzer import evaluate_interview

        ai_report = evaluate_interview(report.get("setup") or {}, report.get("answers") or [])
        ai_report["id"] = save_interview(ai_report, current_user_id())
        return jsonify({"success": True, "report": ai_report})
    except Exception as exc:
        app.logger.exception("Deferred AI evaluation failed")
        return jsonify({"success": False, "error": "AI evaluation is temporarily unavailable because the Gemini API usage limit has been reached. Your interview has been saved and can be evaluated later."}), 429


@app.route("/interview/history")
@login_required
def interview_history():
    return jsonify({"success": True, "interviews": get_interview_history(user_id=current_user_id())})


@app.route("/interview/history/<int:interview_id>")
@login_required
def interview_history_detail(interview_id):
    report = get_interview_report(interview_id, current_user_id())
    if not report:
        return jsonify({"success": False, "error": "Interview report not found."}), 404
    return jsonify({"success": True, "report": report})


@app.route("/interview/history/<int:interview_id>", methods=["DELETE"])
@login_required
def interview_history_delete(interview_id):
    if not delete_interview(interview_id, current_user_id()):
        return jsonify({"success": False, "error": "Interview report not found."}), 404
    return jsonify({"success": True})


@app.route("/interview/history/<int:interview_id>/report-pdf")
@login_required
def interview_history_pdf(interview_id):
    report = get_interview_report(interview_id, current_user_id())
    if not report:
        return jsonify({"success": False, "error": "Interview report not found."}), 404

    try:
        pdf_file = create_interview_pdf(report)
    except Exception:
        app.logger.exception("Stored interview PDF export failed")
        return jsonify({"success": False, "error": "Could not export the stored interview report PDF."}), 400

    setup = report.get("setup") or {}
    safe_name = secure_filename(f"{setup.get('candidate_name') or 'candidate'}-interview-report")
    return send_file(
        pdf_file,
        as_attachment=True,
        download_name=f"{safe_name}.pdf",
        mimetype="application/pdf",
    )


@app.route("/interview/report-pdf", methods=["POST"])
@login_required
def interview_report_pdf():
    payload = request.get_json(silent=True) or {}
    report = payload.get("report") if isinstance(payload.get("report"), dict) else payload
    try:
        pdf_file = create_interview_pdf(report if isinstance(report, dict) else {})
    except Exception:
        app.logger.exception("Interview PDF export failed")
        return jsonify({"success": False, "error": "Could not export the interview report PDF."}), 400

    return send_file(
        pdf_file,
        as_attachment=True,
        download_name="resumeiq-plus-interview-report.pdf",
        mimetype="application/pdf",
    )


@app.route("/aptitude/start", methods=["POST"])
@login_required
def aptitude_start():
    payload = request.get_json(silent=True) or {}
    try:
        question_count = int(payload.get("question_count") or 10)
    except (TypeError, ValueError):
        question_count = 10
    setup = {
        "candidate_name": str(payload.get("candidate_name") or "Candidate").strip()[:120],
        "category": str(payload.get("category") or "Mixed").strip()[:80],
        "difficulty": str(payload.get("difficulty") or "Medium").strip()[:80],
        "question_count": max(1, min(question_count, 30)),
    }
    try:
        from utils.question_bank import get_aptitude_questions

        questions = get_aptitude_questions(setup)
        offline = True
        message = LOCAL_APTITUDE_START_MESSAGE
        public_questions = [{key: value for key, value in item.items() if key != "correct_answer"} for item in questions]
        return jsonify({"success": True, "setup": setup, "questions": public_questions, "answer_key": questions, "offline": offline, "message": message})
    except ValueError as exc:
        app.logger.warning("Aptitude generation failed: %s", exc, exc_info=True)
        return jsonify({"success": False, "error": "Could not start the aptitude test. Please check the setup details and try again."}), 400
    except Exception:
        app.logger.exception("Aptitude setup failed")
        return jsonify({"success": False, "error": "Could not start the aptitude test right now. Please try again."}), 500


@app.route("/aptitude/evaluate", methods=["POST"])
@login_required
def aptitude_evaluate():
    payload = request.get_json(silent=True) or {}
    setup = payload.get("setup") if isinstance(payload.get("setup"), dict) else {}
    questions = payload.get("questions") if isinstance(payload.get("questions"), list) else []
    answers = payload.get("answers") if isinstance(payload.get("answers"), dict) else {}
    try:
        from utils.analyzer import evaluate_aptitude

        result = evaluate_aptitude(setup, questions, answers, payload.get("started_at"), payload.get("finished_at"))
        result["id"] = save_aptitude_test(result, current_user_id())
        return jsonify({"success": True, "result": result})
    except Exception:
        app.logger.exception("Aptitude evaluation failed")
        return jsonify({"success": False, "error": "Could not evaluate aptitude test."}), 400


@app.route("/aptitude/history/<int:test_id>/report-pdf")
@login_required
def aptitude_history_pdf(test_id):
    result = get_aptitude_test(test_id, current_user_id())
    if not result:
        return jsonify({"success": False, "error": "Aptitude test not found."}), 404
    try:
        pdf_file = create_aptitude_pdf(result)
    except Exception:
        app.logger.exception("Aptitude PDF export failed")
        return jsonify({"success": False, "error": "Could not export the aptitude report PDF."}), 400
    return send_file(pdf_file, as_attachment=True, download_name="resumeiq-plus-aptitude-report.pdf", mimetype="application/pdf")


@app.route("/aptitude/history/<int:test_id>", methods=["GET", "DELETE"])
@login_required
def aptitude_history_detail(test_id):
    if request.method == "DELETE":
        if not delete_aptitude_test(test_id, current_user_id()):
            return jsonify({"success": False, "error": "Aptitude test not found."}), 404
        return jsonify({"success": True})
    result = get_aptitude_test(test_id, current_user_id())
    if not result:
        return jsonify({"success": False, "error": "Aptitude test not found."}), 404
    return jsonify({"success": True, "result": result})


@app.route("/aptitude/report-pdf", methods=["POST"])
@login_required
def aptitude_report_pdf():
    payload = request.get_json(silent=True) or {}
    result = payload.get("result") if isinstance(payload.get("result"), dict) else payload
    try:
        pdf_file = create_aptitude_pdf(result if isinstance(result, dict) else {})
    except Exception:
        app.logger.exception("Aptitude PDF export failed")
        return jsonify({"success": False, "error": "Could not export the aptitude report PDF."}), 400
    return send_file(pdf_file, as_attachment=True, download_name="resumeiq-plus-aptitude-report.pdf", mimetype="application/pdf")


@app.route("/resume-pdf", methods=["POST"])
@login_required
def resume_pdf():
    payload = request.get_json(silent=True) or {}
    resume = payload.get("resume") if isinstance(payload.get("resume"), dict) else payload

    if not isinstance(resume, dict):
        return jsonify({"success": False, "error": "Invalid resume data."}), 400

    try:
        save_generated_resume(resume, current_user_id())
        pdf_file = create_resume_pdf(resume)
    except Exception:
        app.logger.exception("Resume PDF generation failed")
        return jsonify({"success": False, "error": "Could not generate the resume PDF."}), 400

    safe_name = secure_filename((resume.get("personal_details") or {}).get("full_name") or "resume")
    return send_file(
        pdf_file,
        as_attachment=True,
        download_name=f"{safe_name or 'resume'}-resume.pdf",
        mimetype="application/pdf",
    )


@app.route("/resume-docx", methods=["POST"])
@login_required
def resume_docx():
    payload = request.get_json(silent=True) or {}
    resume = payload.get("resume") if isinstance(payload.get("resume"), dict) else payload

    if not isinstance(resume, dict):
        return jsonify({"success": False, "error": "Invalid resume data."}), 400

    try:
        save_generated_resume(resume, current_user_id())
        docx_file = create_resume_docx(resume)
    except ValueError as exc:
        return jsonify({"success": False, "error": str(exc)}), 400
    except Exception:
        app.logger.exception("Resume DOCX generation failed")
        return jsonify({"success": False, "error": "Could not generate the resume DOCX."}), 400

    safe_name = secure_filename((resume.get("personal_details") or {}).get("full_name") or "resume")
    return send_file(
        docx_file,
        as_attachment=True,
        download_name=f"{safe_name or 'resume'}-resume.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify({"success": False, "error": "Uploaded file must be smaller than 8 MB."}), 413


@app.errorhandler(HTTPException)
def handle_http_exception(error):
    if wants_json_response() or request.path.startswith("/auth/"):
        message = error.description if error.code and error.code < 500 else "Server unavailable. Please try again shortly."
        if error.code == 404:
            message = "We could not find that page or action."
        elif error.code == 405:
            message = "That action is not available from here. Please refresh and try again."
        return jsonify({"success": False, "error": message}), error.code
    return error


@app.errorhandler(Exception)
def handle_unexpected_exception(error):
    app.logger.exception("Unhandled application error")
    if wants_json_response() or request.path.startswith("/auth/"):
        return jsonify({"success": False, "error": "Server unavailable. Please try again shortly."}), 500
    raise error


@app.route("/api/me")
@login_required
def api_me():
    user_id = current_user_id()
    return jsonify({"success": True, "user": get_user_by_id(user_id), "profile": get_profile(user_id), "settings": get_settings(user_id)})


@app.route("/api/profile", methods=["GET", "POST"])
@login_required
def api_profile():
    user_id = current_user_id()
    if request.method == "GET":
        return jsonify({"success": True, "profile": get_profile(user_id)})
    try:
        profile = update_profile(user_id, request.get_json(silent=True) or request.form)
        return jsonify({"success": True, "profile": profile})
    except Exception as exc:
        app.logger.warning("Profile update failed: %s", exc, exc_info=True)
        return jsonify({"success": False, "error": "Could not update profile. Check the email and fields."}), 400


@app.route("/api/profile/image", methods=["POST"])
@login_required
def api_profile_image():
    image = request.files.get("profile_image")
    if not image or not image.filename:
        return jsonify({"success": False, "error": "Choose a profile image to upload."}), 400
    if not allowed_profile_image(image.filename):
        return jsonify({"success": False, "error": "Profile image must be a JPG, PNG, GIF, or WEBP file."}), 400
    if image.content_length and image.content_length > 2 * 1024 * 1024:
        return jsonify({"success": False, "error": "Profile image must be smaller than 2 MB."}), 400
    user_id = current_user_id()
    extension = secure_filename(image.filename).rsplit(".", 1)[1].lower()
    filename = f"profile-{user_id}-{uuid.uuid4().hex}.{extension}"
    image.save(app.config["UPLOAD_FOLDER"] / filename)
    profile = update_profile(user_id, {"profile_picture": url_for("uploaded_file", filename=filename)})
    return jsonify({"success": True, "message": "Profile image updated.", "profile": profile})


@app.route("/api/onboarding", methods=["POST"])
@login_required
def api_onboarding():
    profile = complete_onboarding(current_user_id(), request.get_json(silent=True) or {})
    return jsonify({"success": True, "profile": profile})


@app.route("/api/settings", methods=["GET", "POST"])
@login_required
def api_settings():
    if request.method == "GET":
        return jsonify({"success": True, "settings": get_settings(current_user_id())})
    settings = update_settings(current_user_id(), request.get_json(silent=True) or {})
    return jsonify({"success": True, "settings": settings})


@app.route("/api/resumes")
@login_required
def api_resumes():
    query = request.args.get("q", "")
    status = request.args.get("status", "active")
    favorite = request.args.get("favorite")
    favorite_flag = None if favorite is None else favorite in {"1", "true", "yes"}
    return jsonify({"success": True, "resumes": list_resumes(current_user_id(), query=query, status=status, favorite=favorite_flag)})


@app.route("/api/resumes/<int:resume_id>/<action>", methods=["POST"])
@login_required
def api_resume_action(resume_id, action):
    payload = request.get_json(silent=True) or {}
    updated = update_resume_state(current_user_id(), resume_id, action, payload.get("title"))
    if not updated:
        return jsonify({"success": False, "error": "Resume not found or action unavailable."}), 404
    return jsonify({"success": True, "resume": updated})


@app.route("/api/templates")
@login_required
def api_templates():
    templates = [
        {"name": "Modern", "category": "Modern", "features": ["Two column", "Impact summary", "ATS readable"]},
        {"name": "Corporate", "category": "Corporate", "features": ["Formal layout", "Leadership-ready", "Clean sections"]},
        {"name": "Developer", "category": "Developer", "features": ["Projects", "Tech stack", "GitHub-ready"]},
        {"name": "Student", "category": "Student", "features": ["Education-first", "Projects", "Certifications"]},
        {"name": "Minimal", "category": "Minimal", "features": ["Single column", "Dense content", "Parser friendly"]},
        {"name": "Executive", "category": "Executive", "features": ["Summary", "Strategy", "Metrics"]},
        {"name": "Creative", "category": "Creative", "features": ["Portfolio links", "Accent color", "Readable flair"]},
        {"name": "ATS Friendly", "category": "ATS Friendly", "features": ["No tables", "Keyword structure", "Plain hierarchy"]},
        {"name": "Software Engineer", "category": "Technology", "features": ["Projects", "Technical skills", "Two column"]},
        {"name": "Data Analyst", "category": "Technology", "features": ["Metrics", "Tools", "Portfolio links"]},
        {"name": "Google Style", "category": "Modern", "features": ["Clean hierarchy", "Impact bullets", "Readable"]},
        {"name": "Microsoft Style", "category": "Corporate", "features": ["Skills-first", "Structured", "Professional"]},
        {"name": "Harvard Style", "category": "Academic", "features": ["Classic", "Single column", "ATS readable"]},
        {"name": "Stanford Style", "category": "Academic", "features": ["Research-ready", "Minimal", "Focused"]},
        {"name": "Two Column", "category": "Layout", "features": ["Compact", "Modern", "Balanced"]},
        {"name": "Single Column", "category": "Layout", "features": ["Parser friendly", "Focused", "Simple"]},
        {"name": "Elegant", "category": "Creative", "features": ["Refined", "Accent color", "Readable"]},
        {"name": "Classic", "category": "Traditional", "features": ["Conservative", "Clear", "Recruiter friendly"]},
        {"name": "Consulting", "category": "Business", "features": ["Results-driven", "Metrics", "Executive"]},
        {"name": "Startup", "category": "Modern", "features": ["Versatile", "Projects", "Impact"]},
    ]
    return jsonify({"success": True, "templates": templates})


@app.route("/api/jobs", methods=["GET", "POST"])
@login_required
def api_jobs():
    user_id = current_user_id()
    if request.method == "GET":
        return jsonify({"success": True, "jobs": list_job_applications(user_id, request.args.get("status", ""))})
    try:
        job_id = save_job_application(user_id, request.get_json(silent=True) or {})
        job = next(item for item in list_job_applications(user_id) if item["id"] == job_id)
        return jsonify({"success": True, "job": job}), 201
    except ValueError as exc:
        return jsonify({"success": False, "error": str(exc)}), 400


@app.route("/api/jobs/<int:job_id>", methods=["PATCH", "DELETE"])
@login_required
def api_job_detail(job_id):
    user_id = current_user_id()
    if request.method == "DELETE":
        if not delete_job_application(user_id, job_id):
            return jsonify({"success": False, "error": "Job application not found."}), 404
        return jsonify({"success": True})
    try:
        job = update_job_application(user_id, job_id, request.get_json(silent=True) or {})
        if not job:
            return jsonify({"success": False, "error": "Job application not found."}), 404
        return jsonify({"success": True, "job": job})
    except ValueError as exc:
        return jsonify({"success": False, "error": str(exc)}), 400


@app.route("/api/notifications")
@login_required
def api_notifications():
    return jsonify({"success": True, "notifications": get_notifications(current_user_id())})


@app.route("/api/search")
@login_required
def api_search():
    return jsonify({"success": True, "results": global_search(current_user_id(), request.args.get("q", ""))})


@app.route("/api/admin")
@login_required
@admin_required
def api_admin():
    return jsonify({"success": True, "admin": get_admin_analytics()})


@app.route("/api/admin/users/<int:user_id>", methods=["DELETE"])
@login_required
@admin_required
def api_admin_delete_user(user_id):
    if user_id == current_user_id():
        return jsonify({"success": False, "error": "You cannot delete your own admin account."}), 400
    if not delete_user(user_id):
        return jsonify({"success": False, "error": "User not found or protected."}), 404
    return jsonify({"success": True})


@app.route("/export-pdf", methods=["POST"])
@login_required
def export_pdf():
    try:
        payload = request.get_json(silent=True) or json.loads(request.form.get("report_data", "{}"))
        pdf_file = create_report_pdf(payload)
    except Exception:
        app.logger.exception("PDF export failed")
        return jsonify({"success": False, "error": "Could not export the report PDF. Please analyze the resume again."}), 400

    return send_file(
        pdf_file,
        as_attachment=True,
        download_name="resumeiq-plus-report.pdf",
        mimetype="application/pdf",
    )


if __name__ == "__main__":
    app.run(debug=os.getenv("FLASK_DEBUG") == "1")
